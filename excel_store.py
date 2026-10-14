"""Synchronous Excel log helpers (call via asyncio.to_thread from async code).

Workbook: data/media_log.xlsx, sheet "media_log".
One writer at a time via a threading lock. Saves go through a temp file +
atomic replace, with retries for the Windows "file open in Excel" case.
"""
import os
import threading
import time

from openpyxl import Workbook, load_workbook

HEADERS = ["timestamp", "original_name", "stored_name", "size_bytes",
           "mime_type", "mode", "url", "message_id", "status", "error"]

SHEET_NAME = "media_log"

_lock = threading.Lock()


def ensure_workbook(path: str) -> None:
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    if os.path.exists(path):
        return
    with _lock:
        if os.path.exists(path):
            return
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        ws.append(HEADERS)
        _save(wb, path)


def _save(wb, path: str, attempts: int = 5, delay: float = 1.0) -> None:
    tmp = path + ".tmp"
    last_err = None
    for _ in range(attempts):
        try:
            wb.save(tmp)
            os.replace(tmp, path)
            return
        except PermissionError as exc:
            last_err = exc
            try:
                if os.path.exists(tmp):
                    os.remove(tmp)
            except OSError:
                pass
            time.sleep(delay)
    raise PermissionError(
        f"Cannot save {path} (is it open in Excel?). Close it and retry."
    ) from last_err


def _get_ws(wb):
    if SHEET_NAME in wb.sheetnames:
        return wb[SHEET_NAME]
    ws = wb.active
    ws.title = SHEET_NAME
    return ws


def append_row(path: str, record: dict) -> int:
    """Append one log row. Returns the Excel row number (for later updates)."""
    ensure_workbook(path)
    with _lock:
        wb = load_workbook(path)
        try:
            ws = _get_ws(wb)
            if ws.max_row == 1 and ws.cell(row=1, column=1).value != HEADERS[0]:
                ws.delete_rows(1, ws.max_row)
                ws.append(HEADERS)
            ws.append([record.get(h, "") for h in HEADERS])
            row_idx = ws.max_row
            _save(wb, path)
        finally:
            wb.close()
        return row_idx


def update_row(path: str, row_idx: int, updates: dict) -> None:
    """Update cells of an existing row by header name."""
    ensure_workbook(path)
    with _lock:
        wb = load_workbook(path)
        try:
            ws = _get_ws(wb)
            col_of = {HEADERS[i]: i + 1 for i in range(len(HEADERS))}
            for key, value in updates.items():
                if key in col_of:
                    ws.cell(row=row_idx, column=col_of[key], value=value)
            _save(wb, path)
        finally:
            wb.close()


def read_all(path: str, limit: int = 100) -> list:
    """Return the last `limit` rows as dicts (oldest first)."""
    ensure_workbook(path)
    with _lock:
        wb = load_workbook(path, read_only=True)
        try:
            ws = _get_ws(wb)
            rows = list(ws.iter_rows(values_only=True))
        finally:
            wb.close()
    if not rows:
        return []
    header = [c for c in rows[0]]
    out = []
    for r in rows[1:]:
        out.append({header[i]: (r[i] if i < len(r) else "") for i in range(len(header))})
    return out[-limit:]
